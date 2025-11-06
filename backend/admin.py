from sqlalchemy.orm import Session
import models_db
from passlib.context import CryptContext
from database_connect import SessionLocal
import secrets
import string
import openpyxl
import os

db = SessionLocal()


def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


def add_presentations_to_database():
    file_path = os.path.join(os.getcwd(), "data/lectures_and_posters.xlsx")
    exel_users_file = openpyxl.load_workbook(file_path)
    sheets = exel_users_file.sheetnames
    lectures, posters = exel_users_file[sheets[0]], exel_users_file[sheets[1]]

    # Lectures
    for row in range(2, lectures.max_row + 1):

        lecture_code = "".join(
            secrets.choice(string.ascii_uppercase + string.digits) for _ in range(4)
        )

        lectures.cell(row=row, column=6, value=lecture_code)

        # nie walnijcie sie z kolejnoscia kolumn w exelu
        speaker_name = lectures.cell(row=row, column=1).value
        lecture_category = lectures.cell(row=row, column=2).value
        lecture_name = lectures.cell(row=row, column=3).value

        new_lecture = models_db.Lectures(
            lecture_name=lecture_name,
            speaker_name=speaker_name,
            lecture_category=lecture_category,
            lecture_code=lecture_code,
        )

        db.add(new_lecture)

    for row in range(2, posters.max_row + 1):

        poster_author = posters.cell(row=row, column=1).value
        poster_name = posters.cell(row=row, column=2).value

        new_poster = models_db.Posters(
            poster_name=poster_name,
            poster_author=poster_author,
        )

        db.add(new_poster)

    db.commit()
    db.close()


pwd_context = CryptContext(schemes=["bcrypt"], deprecated="auto")


def add_users_to_database():
    file_path = os.path.join(os.getcwd(), "data/participants.xlsx")
    exel_users_file = openpyxl.load_workbook(file_path)
    worksheet = exel_users_file.active

    for row in range(2, worksheet.max_row + 1):

        user_login = (
            (
                worksheet.cell(row=row, column=1).value
                + worksheet.cell(row=row, column=2).value
            )
            .replace(" ", "")
            .lower()
        )
        password_length = 5
        password = secrets.token_urlsafe(password_length)
        hashed_password = pwd_context.hash(password)

        worksheet.cell(row=row, column=4, value=user_login)
        worksheet.cell(row=row, column=5, value=password)

        new_user = models_db.Users(user_login=user_login, user_password=hashed_password)
        db.add(new_user)

    exel_users_file.save(file_path)

    db.commit()
    db.close()


def update_all_lecture_points():
    lectures = db.query(models_db.Lectures).all()

    if not lectures:
        print("No lectures found")
        return

    for lecture in lectures:
        votes = (
            db.query(models_db.Votes)
            .filter(models_db.Votes.lecture_id == lecture.lecture_id)
            .all()
        )
        merytoryka_total, forma_total, merytoryka_count, forma_count = 0.0, 0.0, 0, 0

        for vote in votes:
            if vote.merytoryka_points is not None:
                merytoryka_total += vote.merytoryka_points
                merytoryka_count += 1
            if vote.forma_points is not None:
                forma_total += vote.forma_points
                forma_count += 1

        merytoryka_average = (
            merytoryka_total / merytoryka_count if merytoryka_count else 0
        )
        forma_average = forma_total / forma_count if forma_count else 0
        weighted_average = 0.6 * merytoryka_average + 0.4 * forma_average

        lecture.sum_points = weighted_average

    db.commit()
    print("Lecture points updated successfully for all lectures")


def update_posters_points():
    posters = db.query(models_db.Posters).all()

    if not posters:
        print("No posters found")
        return

    for poster in posters:
        votes = (
            db.query(models_db.Votes_posters)
            .filter(models_db.Votes_posters.poster_id == poster.poster_id)
            .all()
        )
        merytoryka_total, estetyka_total, merytoryka_count, estetyka_count = (
            0.0,
            0.0,
            0,
            0,
        )

        for vote in votes:
            if vote.merytoryka_points is not None:
                merytoryka_total += vote.merytoryka_points
                merytoryka_count += 1
            if vote.estetyka_points is not None:
                estetyka_total += vote.estetyka_points
                estetyka_count += 1

        merytoryka_average = (
            merytoryka_total / merytoryka_count if merytoryka_count else 0
        )
        estetyka_average = estetyka_total / estetyka_count if estetyka_count else 0
        weighted_average = 0.8 * merytoryka_average + 0.2 * estetyka_average

        poster.sum_points = weighted_average

    db.commit()
    print("Posters points updated successfully for all posters")


# add_users_to_database()
# add_presentations_to_database()

# update_all_lecture_points()
# update_posters_points()
